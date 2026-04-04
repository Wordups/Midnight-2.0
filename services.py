import os
import tempfile
from datetime import datetime
from docx import Document
from groq import Groq
from hps_policy_migration_builder import build_policy_document

LOCAL_GROQ_API_KEY = ""

EXTRACTION_PROMPT = """
You are a policy migration specialist.

Your task is to read the attached legacy policy document and extract ALL content
into the exact Python dictionary structure below.

STRICT RULES:
- Do NOT summarize, rewrite, or remove content
- Preserve the source wording as closely as possible
- Fix only minor spacing / punctuation / obvious grammar defects where needed
- Map all content into the correct field
- If content does not fit perfectly, place it in the most logical field
- For procedure items, classify each entry using exactly one type:
  "para"           = standalone paragraph
  "heading"        = bold underlined subsection title
  "bullet"         = first-level bullet
  "sub-bullet"     = second-level bullet
  "bold_intro"     = paragraph that starts with a bold label; use keys "bold" and "rest"
  "bold_intro_semi"= same as bold_intro but "rest" contains semicolons
  "empty"          = blank spacer line

Return ONLY a valid Python dictionary assignment. No explanation. No markdown.
Start your response with:
POLICY_DATA = {
and end with the closing brace.

POLICY_DATA = {
    "policy_name": "",
    "policy_number": "",
    "version": "",
    "grc_id": "",
    "supersedes": "",
    "effective_date": "",
    "last_reviewed": "",
    "last_revised": "",
    "custodians": "",
    "owner_name": "",
    "owner_title": "",
    "approver_name": "",
    "approver_title": "",
    "date_signed": "",
    "date_approved": "",
    "applicable_to": {
        "hps_inc": False,
        "agency": True,
        "corporate": True,
        "govt_affairs": False,
        "legal_review": False
    },
    "policy_types": {
        "carrier_specific": False,
        "cross_carrier": False,
        "global": False,
        "on_off_hix": False
    },
    "line_of_business": {
        "all_lobs": True,
        "specific_lob": "",
        "specific_lob_checked": False
    },
    "purpose": "",
    "definitions": {},
    "policy_statement": "",
    "procedures": [],
    "related_policies": [],
    "citations": [],
    "revision_history": []
}

HERE IS THE LEGACY POLICY DOCUMENT:
"""

def get_api_key():
    return os.getenv("GROQ_API_KEY", "") or LOCAL_GROQ_API_KEY

def parse_policy_data(raw_output: str):
    if "POLICY_DATA = {" in raw_output:
        raw_output = raw_output[raw_output.index("POLICY_DATA = {"):]
    raw_output = (
        raw_output.replace("\u201c", '"')
        .replace("\u201d", '"')
        .replace("\u2019", "'")
    )
    namespace = {}
    exec(raw_output, {}, namespace)
    return namespace.get("POLICY_DATA")

def extract_text_from_docx(file_obj):
    doc = Document(file_obj)
    lines = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            lines.append(text)

    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                ct = " ".join(
                    para.text.strip() for para in cell.paragraphs if para.text.strip()
                ).strip()
                if ct:
                    row_text.append(ct)
            if row_text:
                lines.append(" | ".join(row_text))

    return "\n".join(lines)

def get_uploaded_text(filename, file_bytes):
    from io import BytesIO
    if filename.lower().endswith(".docx"):
        return extract_text_from_docx(BytesIO(file_bytes))
    return file_bytes.decode("utf-8", errors="ignore")

def save_logo_bytes(filename, file_bytes):
    safe_name = os.path.basename(filename)
    path = os.path.join(tempfile.gettempdir(), safe_name)
    with open(path, "wb") as f:
        f.write(file_bytes)
    return path

def run_llm_transform(source_text, template_name):
    api_key = get_api_key()
    if not api_key:
        raise RuntimeError("Missing GROQ_API_KEY")

    client = Groq(api_key=api_key)

    response = client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        messages=[{"role": "user", "content": EXTRACTION_PROMPT + "\n\n" + source_text}],
        temperature=0.1,
        max_tokens=8000,
    )

    raw = response.choices[0].message.content.strip()
    data = parse_policy_data(raw)

    if not data:
        raise RuntimeError("Model response could not be parsed into POLICY_DATA.")

    data["template_name"] = template_name
    return data

def build_output_doc(policy_data, logo_path=None):
    name = policy_data.get("policy_name", "Policy")
    number = policy_data.get("policy_number", "SEC-P")
    version = policy_data.get("version", "V1.0")
    filename = f"{number} {name} {version}-NEW.docx"

    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        path = tmp.name

    build_policy_document(policy_data, path, logo_path=logo_path)

    with open(path, "rb") as f:
        data = f.read()

    os.unlink(path)
    return filename, data
